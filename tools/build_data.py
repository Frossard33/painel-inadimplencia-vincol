import openpyxl,sys,re,json,datetime,collections
src,out=sys.argv[1],sys.argv[2]
MASK=True
wb=openpyxl.load_workbook(src,data_only=True)
def find(prefix):
    return [w for w in wb if w.title.startswith(prefix)]
TABS=[("Dív","Dívida Antiga"),("Judicial","Judicial"),("Org","Órgão Público"),("Frota","Frota"),("Auditoria","Auditoria"),("Permuta","Permuta"),("Dep","Depósito"),("Pneustore","Pneustore")]
ENT=re.compile(r"\b(LTDA|ME|EPP|EIRELI|S/?A|S\.A\.?|COMERCIO|COMÉRCIO|TRANSPORTES?|ASSOCIACAO|ASSOCIAÇÃO|FUNDO|INSTITUTO|PREFEITURA|MUNICIPAL|SERV|SERVICOS|SERVIÇOS|EMPRESA|CIA|CIA\.|LOGISTICA|AUTO|PECAS|PEÇAS|PNEUS?|BORRACHARIA|VIDROSCAR|CONCRETO|SOLUCOES|SOLUÇÕES|CORPORATIVAS|MERCADO|AUTOMOTIVA|AUTOMOTORES|REDE|CENTRO|CARGO|GESSO|SAAE|COOPERATIVA|CONSTRUCAO|CONSTRUÇÃO|MATERIAL|INDUSTRIA|INDÚSTRIA|RECAUCHUTADORA|TRANSP|LOCADORA|MOTORS|SERVICE|CENTER)\b",re.I)
def mask(n):
    n=(n or "").strip()
    if not MASK or ENT.search(n): return n
    n=re.sub(r"^[\d\.\-/ ]+(?=[A-Za-z])","",n)  # remove CNPJ prefix
    p=n.split()
    if len(p)<=1: return n
    return p[0]+" "+" ".join(x[0]+"." for x in p[1:] if x[0].isalpha())
def d(v):
    if v is None or v=="": return None
    if isinstance(v,datetime.datetime): return v.date().isoformat()
    if isinstance(v,datetime.date): return v.isoformat()
    s=str(v).strip()
    m=re.match(r"(\d\d)/(\d\d)/(\d{4})",s)
    if m: return f"{m[3]}-{m[2]}-{m[1]}"
    return s[:10]
def rows(ws):
    hdr=[c.value for c in ws[1]][:18]
    for r in ws.iter_rows(min_row=2,max_col=18,values_only=True):
        if r[0] is None: continue
        yield dict(zip(hdr,r))
def num(v):
    try: return round(float(v),2)
    except: return 0.0
FM={'1':'01','2':'02','3':'03','4':'04','6':'06'}
def mk(r,cat,tab):
    h=list(r.keys())
    g=lambda sub:[r[k] for k in h if k and sub in k][0] if any(k and sub in k for k in h) else None
    f=str(r[h[0]]).strip()
    paid=r.get('Data de Pagamento') or r.get('Data da baixa')
    return dict(f=FM.get(f,'OUT'),fo=f,c=cat,t=tab,
      d=str(g('Pagar')).strip().upper(),tit=str(g('Titulo C')).strip(),par=g('Parcela'),cod=str(g('Fornecedor C')).strip(),
      cli=mask(g('Fornecedor Nome')),v=d(g('Vencimento')),val=num(g('T')) if False else num(r.get('Título Valor') or r.get('T\ufffdtulo Valor')),
      sal=num(r.get('Título Saldo') or r.get('T\ufffdtulo Saldo')),
      pg=d(r.get('Data de Pagamento')),vp=num(r.get('Valor')) if r.get('Data de Pagamento') else 0)
EOF=None
out_rows=[]
for pre,cat in TABS:
    ws=[w for w in wb if w.title.startswith(pre)]
    assert len(ws)==1,(pre,[w.title for w in ws])
    for r in rows(ws[0]):
        c2=cat
        if pre=="Dív": c2="Cobrança Extra" if str(r.get("STATUS")).upper().startswith("COBRANCA EXTRA") else "Dívida Antiga"
        out_rows.append(mk(r,c2,ws[0].title))
for ws in find("Cobran"):
    for r in rows(ws): out_rows.append(mk(r,"Boletos",ws.title))
for r in rows(wb['Base de dados']):
    if r.get('STATUS')=='TESOURARIA': out_rows.append(mk(r,"Tesouraria","Base de dados"))
pagar=[]
for ws in find("Contas a Pagar"):
    for r in rows(ws): pagar.append(mk(r,"Contas a Pagar",ws.title))
print("pagar",len(pagar),round(sum(x["sal"] for x in pagar)))
pend=[]
for ws in wb:
    T=ws.title; cat=None
    if "Vanessa" in T: cat="Conciliação Vanessa"
    elif "Marla" in T: cat="Conciliação Marla"
    elif T.startswith("Pendencias"): cat="Pendências Diversas"
    elif "Koncilli" in T: cat="Mercado Livre Koncilli"
    elif T.startswith("Mercado Livre") and "Conta" in T: cat="Mercado Livre Conta"
    if cat:
        n0=len(pend)
        for r in rows(ws): pend.append(mk(r,cat,T))
        print(cat,len(pend)-n0,round(sum(x["sal"] for x in pend[n0:])))
seen=collections.Counter()
for x in out_rows: seen[(x['fo'],x['tit'],x['par'],x['v'],x['sal'])]+=1
dups=sum(1 for k,v in seen.items() if v>1)
print("rows",len(out_rows),"dup keys",dups)
for c,n in collections.Counter(x['c'] for x in out_rows).items(): print(c,n,round(sum(x['sal'] for x in out_rows if x['c']==c)))
print(collections.Counter(x['f'] for x in out_rows))
print("missing venc",sum(1 for x in out_rows if not x['v']),"zero sal",sum(1 for x in out_rows if x['sal']<=0))
open(out,'w',encoding='utf-8').write("window.DADOS="+json.dumps({"gerado":"2026-10-06","titulos":out_rows,"pagar":pagar,"pend":pend},ensure_ascii=False,separators=(',',':'))+";\n")
